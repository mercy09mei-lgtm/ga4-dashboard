import streamlit as st
import os
import re
import csv
import io
import shutil
import tempfile
import openpyxl
from datetime import datetime
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

# 页面基础配置
st.set_page_config(page_title="GA4 自动报表与看板系统", layout="wide", page_icon="📊")

# ================= 业务映射规则 =================
DOWNLOAD_RULES = [
    (r"sidify-music-converter(-arm)?\.exe", "spotify win"),
    (r"sidify-music-converter(-arm)?\.dmg", "spotify mac"),
    (r"sidify.*\.apk", "spotify安卓"),
    (r"sidify-music-converter-free", "SD Free"),
    (r"sidify-all-in-one(-arm)?\.exe", "musicOne win"),
    (r"sidify-all-in-one(-arm)?\.dmg", "musicOne mac"),
    (r"(amazon-music-converter|sidify-amc)(-arm)?\.exe", "amazC win"),
    (r"amazon-music-converter(-arm)?\.dmg", "amazC mac"),
    (r"tidal-music-converter(-arm)?\.exe", "Tidal win"),
    (r"tidal-music-converter(-arm)?\.dmg", "Tidal mac"),
    (r"apple-music-converter(-arm)?\.exe", "Apple win"),
    (r"apple-music-converter(-arm)?\.dmg", "Apple mac"),
    (r"youtube-music-converter(-arm)?\.exe", "YouTube win"),
    (r"youtube-music-converter(-arm)?\.dmg", "YouTube mac"),
    (r"deezer-music-converter(-arm)?\.exe", "Deezer win"),
    (r"deezer-music-converter(-arm)?\.dmg", "Deezer mac"),
]

PRODUCT_MAP = {
    "sp": "spotify", "spotify": "spotify",
    "all-in-one": "musicOne", "all_in_one": "musicOne", "musicone": "musicOne",
    "az": "amazC", "amazon": "amazC",
    "tidal": "Tidal", "td": "Tidal",
    "apc": "Apple", "apple": "Apple",
    "youtube": "YouTube", "yt": "YouTube",
    "deezer": "Deezer", "dz": "Deezer",
}

PLAN_MAP = {
    "lifetime": "lifetime",
    "one_year": "12M", "12m": "12M", "1_year": "12M",
    "3m": "3M", "three_month": "3M",
    "1m": "1M", "one_month": "1M",
}

TEMPLATE_FILE = os.path.join(os.path.dirname(__file__), "template.xlsx")

# ================= 数据解析函数 =================
def analyze_date_info_from_text(text):
    for line in text.splitlines()[:15]:
        match = re.search(r'#\s*(\d{8})-(\d{8})\s+(\d{8})-(\d{8})', line)
        if match:
            c_start, c_end, p_start, p_end = match.groups()
            d_c_start = datetime.strptime(c_start, "%Y%m%d")
            d_c_end = datetime.strptime(c_end, "%Y%m%d")
            d_p_start = datetime.strptime(p_start, "%Y%m%d")
            d_p_end = datetime.strptime(p_end, "%Y%m%d")
            days_span = (d_c_end - d_c_start).days + 1

            if days_span >= 20:
                report_type = "month"
                file_tag = f"{d_c_start.strftime('%Y%m')}"
                cur_label = f"{d_c_start.month}月"
                prev_label = f"{d_p_start.month}月"
            else:
                report_type = "week"
                file_tag = f"{c_start}-{c_end}"
                cur_label = f"{d_c_start.month}.{d_c_start.day}-{d_c_end.month}.{d_c_end.day}"
                prev_label = f"{d_p_start.month}.{d_p_start.day}-{d_p_end.month}.{d_p_end.day}"

            return {"type": report_type, "days": days_span, "file_tag": file_tag, "cur_label": cur_label, "prev_label": prev_label}
    return {"type": "unknown", "days": 7, "file_tag": "UNKNOWN", "cur_label": "本期", "prev_label": "上期"}

def parse_site_name_from_traffic(text):
    for line in text.splitlines()[:10]:
        match = re.search(r'https?://(?:www\.)?([^/\s\-]+)', line)
        if match:
            raw_host = match.group(1)
            parts = raw_host.split('.')
            if len(parts) >= 2:
                return f"{parts[0].capitalize()}.{parts[1]}"
            return raw_host.capitalize()
    return "Sidify.de"

def parse_download_csv_content(text):
    rows = list(csv.reader(text.splitlines()))
    desktop_idx, mobile_idx, tablet_idx, totals_idx = 2, 3, None, -1

    for row in rows[:15]:
        if any("设备类别" in str(cell) for cell in row):
            for idx, val in enumerate(row):
                v_clean = val.strip().lower()
                if v_clean == "desktop": desktop_idx = idx
                elif v_clean == "mobile": mobile_idx = idx
                elif v_clean == "tablet": tablet_idx = idx
                elif v_clean in ["totals", "合计"]: totals_idx = idx
            break

    def safe_parse_val(raw_str):
        if not raw_str or raw_str.strip() in ['', '-', 'None']: return 0.0
        return float(raw_str.replace(',', '').strip())

    cur_period_label, prev_period_label = None, None
    current_key = None
    data = {}
    ga4_totals = {"cur": 0, "prev": 0}

    for row in rows:
        if not row: continue
        first_col = row[0].strip()
        if first_col.startswith('http') or any(first_col.endswith(ext) for ext in ['.exe', '.dmg', '.apk']):
            current_key = first_col

        if len(row) > 1 and " - " in row[1] and "变动百分比" not in row[1]:
            label = row[1].strip()
            if not cur_period_label: cur_period_label = label
            elif not prev_period_label and label != cur_period_label: prev_period_label = label

            if not current_key and first_col == "":
                tot_val = safe_parse_val(row[totals_idx]) if len(row) > abs(totals_idx) else 0.0
                if label == cur_period_label:
                    ga4_totals["cur"] = int(tot_val)
                else:
                    ga4_totals["prev"] = int(tot_val)

            if current_key:
                target_col = None
                for pat, col_name in DOWNLOAD_RULES:
                    if re.search(pat, current_key, re.IGNORECASE):
                        target_col = col_name
                        break
                if target_col:
                    if target_col == "spotify安卓":
                        m_val = safe_parse_val(row[mobile_idx]) if len(row) > mobile_idx else 0.0
                        t_val = safe_parse_val(row[tablet_idx]) if tablet_idx is not None and len(row) > tablet_idx else 0.0
                        val = m_val + t_val
                    else:
                        val = safe_parse_val(row[desktop_idx]) if len(row) > desktop_idx else 0.0

                    if target_col not in data:
                        data[target_col] = {"cur": 0, "prev": 0}
                    if label == cur_period_label:
                        data[target_col]["cur"] += int(val)
                    else:
                        data[target_col]["prev"] += int(val)

    return data, ga4_totals

def parse_buy_csv_content(text):
    rows = list(csv.reader(text.splitlines()))
    cur_period_label, prev_period_label = None, None
    current_key = None
    data = {"lifetime": {}, "12M": {}, "3M": {}, "1M": {}}
    ga4_totals = {"cur": 0, "prev": 0}

    for row in rows:
        if not row: continue
        first_col = row[0].strip()
        if first_col.startswith("buy_"):
            current_key = first_col

        if len(row) > 1 and " - " in row[1] and "变动百分比" not in row[1]:
            label = row[1].strip()
            raw_v = row[2] if len(row) > 2 else "0"
            val = float(raw_v.replace(',', '')) if raw_v not in ['', '-', 'None'] else 0.0

            if not cur_period_label: cur_period_label = label
            elif not prev_period_label and label != cur_period_label: prev_period_label = label

            if not current_key and first_col == "":
                if label == cur_period_label:
                    ga4_totals["cur"] = int(val)
                else:
                    ga4_totals["prev"] = int(val)

            if current_key:
                parts = current_key.split('_')
                if len(parts) == 3:
                    raw_prod, raw_os, raw_plan = parts[1].lower(), "win", parts[2].lower()
                elif len(parts) >= 4:
                    raw_prod = parts[1].lower()
                    if parts[2].lower() in ["win", "mac", "android"]:
                        raw_os = parts[2].lower()
                        raw_plan = "_".join(parts[3:]).lower()
                    else:
                        raw_os = "win"
                        raw_plan = "_".join(parts[2:]).lower()
                else:
                    continue

                norm_prod = PRODUCT_MAP.get(raw_prod, raw_prod)
                norm_plan = PLAN_MAP.get(raw_plan, raw_plan)
                target_col = f"{norm_prod} {raw_os}"

                if norm_plan in data:
                    if target_col not in data[norm_plan]:
                        data[norm_plan][target_col] = {"cur": 0, "prev": 0}
                    if label == cur_period_label:
                        data[norm_plan][target_col]["cur"] += int(val)
                    else:
                        data[norm_plan][target_col]["prev"] += int(val)

    return data, ga4_totals

def parse_traffic_csv_content(text):
    lines = text.splitlines()
    traffic = {"cur": {}, "prev": {}}
    current_period = None

    for line in lines:
        line_str = line.strip()
        if not line_str: continue

        if line_str.startswith("# 开始日期："):
            if current_period is None: current_period = "cur"
            elif current_period == "cur": current_period = "prev"
            continue

        if line_str.startswith('#') or "带来用户首次互动的主渠道组" in line_str:
            continue

        if current_period in ["cur", "prev"]:
            cols = [c.strip() for c in line_str.split(',')]
            if cols and cols[0]:
                channel = cols[0]
                try: users = int(cols[1].replace(',', ''))
                except: users = 0
                traffic[current_period][channel] = users
    return traffic

def generate_report_excel(template_path, date_info, dl_data, dl_ga4_totals, buy_data, buy_ga4_totals, tf_data):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_name = tmp.name

    shutil.copyfile(template_path, tmp_name)
    wb = openpyxl.load_workbook(tmp_name)
    ws = wb.active

    # 1. 周期标签
    for r in [2, 8, 14, 20, 26]:
        ws.cell(r, 1, date_info["cur_label"])
        ws.cell(r + 1, 1, date_info["prev_label"])

    # 2. 填充下载量（各单品分别填入，最后一列保留 SUM 算术求和）
    for c in range(2, 18):
        h = ws.cell(1, c).value
        if h:
            hc = " ".join(str(h).split())
            match = next((v for k, v in dl_data.items() if " ".join(k.split()) == hc), None)
            ws.cell(2, c, match["cur"] if match else 0)
            ws.cell(3, c, match["prev"] if match else 0)
    # 单品总计公式
    ws.cell(2, 18, "=SUM(B2:Q2)")
    ws.cell(3, 18, "=SUM(B3:Q3)")

    # 3. 填充购买量（各单品分别填入，最后一列保留 SUM 算术求和）
    for plan_key, head_r, cur_r, prev_r in [("lifetime", 7, 8, 9), ("12M", 13, 14, 15), ("3M", 19, 20, 21), ("1M", 25, 26, 27)]:
        cur_plan = buy_data.get(plan_key, {})
        for c in range(2, 16):
            h = ws.cell(head_r, c).value
            if h:
                hc = " ".join(str(h).split())
                match = next((v for k, v in cur_plan.items() if " ".join(k.split()) == hc), None)
                ws.cell(cur_r, c, match["cur"] if match else 0)
                ws.cell(prev_r, c, match["prev"] if match else 0)
        ws.cell(cur_r, 16, f"=SUM(B{cur_r}:O{cur_r})")
        ws.cell(prev_r, 16, f"=SUM(B{prev_r}:O{prev_r})")

    # 4. 流量与公式 (Row 31~39)
    if tf_data:
        days = date_info["days"]
        div = days if date_info["type"] == "month" else 7

        for r in range(33, 40):
            ch = str(ws.cell(r, 1).value or "").strip()
            cur_val = tf_data["cur"].get(ch, 0)
            prev_val = tf_data["prev"].get(ch, 0)
            ws.cell(r, 2, cur_val)
            ws.cell(r, 3, prev_val)
            ws.cell(r, 4, f"=B{r}-C{r}")
            ws.cell(r, 9, f"=IF(C{r}=0, 0, (B{r}-C{r})/C{r})")

        total_cur = sum(tf_data["cur"].values())
        total_prev = sum(tf_data["prev"].values())
        ws.cell(32, 2, total_cur)
        ws.cell(32, 3, total_prev)
        ws.cell(32, 4, "=B32-C32")
        ws.cell(32, 5, f"=B32/{div}")
        ws.cell(32, 6, f"=E32-(C32/{div})")
        ws.cell(32, 7, "=B34+B35")
        ws.cell(32, 8, "=C34+C35")
        ws.cell(32, 9, "=IF(H32=0, 0, (G32-H32)/H32)")
        ws.cell(32, 10, "=B32/30")
        ws.cell(32, 11, "=C32/30")

    # 5. 填入核心概览汇总表格 (Row 41~46)
    # 取 GA4 官方去重后的下载与购买总数
    dl_c_tot = dl_ga4_totals["cur"] if dl_ga4_totals["cur"] > 0 else sum(v["cur"] for v in dl_data.values())
    dl_p_tot = dl_ga4_totals["prev"] if dl_ga4_totals["prev"] > 0 else sum(v["prev"] for v in dl_data.values())

    buy_c_tot = buy_ga4_totals["cur"] if buy_ga4_totals["cur"] > 0 else sum(sum(prod["cur"] for prod in p.values()) for p in buy_data.values())
    buy_p_tot = buy_ga4_totals["prev"] if buy_ga4_totals["prev"] > 0 else sum(sum(prod["prev"] for prod in p.values()) for p in buy_data.values())

    # 表头行 41
    ws.cell(41, 1, "指标")
    ws.cell(41, 2, date_info["cur_label"])
    ws.cell(41, 3, date_info["prev_label"])
    ws.cell(41, 4, "环比")
    ws.cell(41, 5, "状态")

    # 42行: 自然+引荐
    ws.cell(42, 1, "自然+引荐")
    ws.cell(42, 2, "=G32")
    ws.cell(42, 3, "=H32")
    ws.cell(42, 4, "=IF(C42=0, 0, (B42-C42)/C42)")
    ws.cell(42, 5, '=IF(D42>0.005, "上涨", IF(D42<-0.005, "下降", "正常"))')

    # 43行: 下载量 (GA4官方去重总量)
    ws.cell(43, 1, "下载量")
    ws.cell(43, 2, dl_c_tot)
    ws.cell(43, 3, dl_p_tot)
    ws.cell(43, 4, "=IF(C43=0, 0, (B43-C43)/C43)")
    ws.cell(43, 5, '=IF(D43>0.005, "上涨", IF(D43<-0.005, "下降", "正常"))')

    # 44行: 购买事件数 (GA4官方去重总量)
    ws.cell(44, 1, "购买事件数")
    ws.cell(44, 2, buy_c_tot)
    ws.cell(44, 3, buy_p_tot)
    ws.cell(44, 4, "=IF(C44=0, 0, (B44-C44)/C44)")
    ws.cell(44, 5, '=IF(D44>0.005, "上涨", IF(D44<-0.005, "下降", "正常"))')

    # 45行: 下载转化率
    ws.cell(45, 1, "下载转化率")
    ws.cell(45, 2, "=IF(B42=0, 0, B43/B42)")
    ws.cell(45, 3, "=IF(C42=0, 0, C43/C42)")
    ws.cell(45, 4, "=B45-C45")
    ws.cell(45, 5, '=IF(D45>0.0005, "上涨", IF(D45<-0.0005, "下降", "正常"))')

    # 46行: 购买转化率
    ws.cell(46, 1, "购买转化率")
    ws.cell(46, 2, "=IF(B43=0, 0, B44/B43)")
    ws.cell(46, 3, "=IF(C43=0, 0, C44/C43)")
    ws.cell(46, 4, "=B46-C46")
    ws.cell(46, 5, '=IF(D46>0.0005, "上涨", IF(D46<-0.0005, "下降", "正常"))')

    # 设置百分比格式
    for r in [42, 43, 44]:
        ws.cell(r, 4).number_format = '0.00%'
    for r in [45, 46]:
        ws.cell(r, 2).number_format = '0.00%'
        ws.cell(r, 3).number_format = '0.00%'
        ws.cell(r, 4).number_format = '+0.00%;-0.00%;0.00%'

    wb.save(tmp_name)
    with open(tmp_name, "rb") as f:
        bytes_data = f.read()
    os.remove(tmp_name)
    return bytes_data


# ================= 网页 UI 布局 =================
st.title("📊 GA4 数据可视化看板与报表系统")

with st.sidebar:
    st.header("📂 数据上传区")
    uploaded_files = st.file_uploader(
        "上传 GA4 导出的 CSV 文件（支持多选拖拽）：",
        type=["csv"],
        accept_multiple_files=True
    )
    st.markdown("---")
    st.subheader("📑 报表模板设置")
    if os.path.exists(TEMPLATE_FILE):
        st.success("✅ 网站已预置默认 `template.xlsx` 模板")
    else:
        st.error("❌ 网站目录下未检测到 `template.xlsx`")

if uploaded_files:
    dl_text, buy_text, tf_text = None, None, None
    for f in uploaded_files:
        name = f.name.lower()
        content = f.getvalue().decode('utf-8', errors='ignore')
        if "download" in name: dl_text = content
        elif "buy" in name: buy_text = content
        elif "用户获取" in name or "渠道" in name: tf_text = content

    if not dl_text or not buy_text:
        st.warning("请确保上传的文件中包含 `download.csv` 和 `buy.csv`！")
    else:
        date_info = analyze_date_info_from_text(dl_text)
        dl_data, dl_ga4_totals = parse_download_csv_content(dl_text)
        buy_data, buy_ga4_totals = parse_buy_csv_content(buy_text)
        tf_data = parse_traffic_csv_content(tf_text) if tf_text else None

        site_name = parse_site_name_from_traffic(tf_text) if tf_text else "Sidify.de"
        total_traffic_cur = sum(tf_data["cur"].values()) if tf_data else 0
        days_count = date_info["days"] if date_info["type"] == "month" else 7
        daily_ip = int(round(total_traffic_cur / days_count)) if days_count > 0 else 0
        ip_label = "月日均 IP" if date_info["type"] == "month" else "周日均 IP"

        top_c1, top_c2, top_c3 = st.columns([2, 2, 3])
        top_c1.metric("周期类型", f"{date_info['type'].upper()}报")
        top_c2.metric("当前对比周期", f"{date_info['cur_label']} vs {date_info['prev_label']}")

        if os.path.exists(TEMPLATE_FILE):
            excel_bytes = generate_report_excel(TEMPLATE_FILE, date_info, dl_data, dl_ga4_totals, buy_data, buy_ga4_totals, tf_data)
            output_name = f"GA4_Report_{date_info['file_tag']}.xlsx"
            top_c3.download_button(
                label=f"📥 导出并下载标准化报表 ({output_name})",
                data=excel_bytes,
                file_name=output_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

        st.markdown("---")

        # ================= 核心概览指标卡 =================
        st.subheader("📌 核心指标概览看板")

        nat_cur = (tf_data["cur"].get("Organic Search", 0) + tf_data["cur"].get("Referral", 0)) if tf_data else 0
        nat_prev = (tf_data["prev"].get("Organic Search", 0) + tf_data["prev"].get("Referral", 0)) if tf_data else 0

        dl_cur_total = dl_ga4_totals["cur"] if dl_ga4_totals["cur"] > 0 else sum(v["cur"] for v in dl_data.values())
        dl_prev_total = dl_ga4_totals["prev"] if dl_ga4_totals["prev"] > 0 else sum(v["prev"] for v in dl_data.values())

        buy_cur_total = buy_ga4_totals["cur"] if buy_ga4_totals["cur"] > 0 else sum(sum(prod["cur"] for prod in p_dict.values()) for p_dict in buy_data.values())
        buy_prev_total = buy_ga4_totals["prev"] if buy_ga4_totals["prev"] > 0 else sum(sum(prod["prev"] for prod in p_dict.values()) for p_dict in buy_data.values())

        dl_c_rate = (dl_cur_total / nat_cur) if nat_cur > 0 else 0.0
        dl_p_rate = (dl_prev_total / nat_prev) if nat_prev > 0 else 0.0
        buy_c_rate = (buy_cur_total / dl_cur_total) if dl_cur_total > 0 else 0.0
        buy_p_rate = (buy_prev_total / dl_prev_total) if dl_prev_total > 0 else 0.0

        def calc_ratio(c, p):
            if p == 0: return 0.0
            return (c - p) / p

        def get_status(diff):
            if abs(diff) < 0.005: return "正常"
            return "上涨" if diff > 0 else "下降"

        overview_rows = [
            {
                "指标": "自然+引荐",
                f"{date_info['cur_label']}": f"{nat_cur:,}",
                f"{date_info['prev_label']}": f"{nat_prev:,}",
                "环比": f"{calc_ratio(nat_cur, nat_prev)*100:+.2f}%",
                "状态": get_status(calc_ratio(nat_cur, nat_prev))
            },
            {
                "指标": "下载量",
                f"{date_info['cur_label']}": f"{dl_cur_total:,}",
                f"{date_info['prev_label']}": f"{dl_prev_total:,}",
                "环比": f"{calc_ratio(dl_cur_total, dl_prev_total)*100:+.2f}%",
                "状态": get_status(calc_ratio(dl_cur_total, dl_prev_total))
            },
            {
                "指标": "购买事件数",
                f"{date_info['cur_label']}": f"{buy_cur_total:,}",
                f"{date_info['prev_label']}": f"{buy_prev_total:,}",
                "环比": f"{calc_ratio(buy_cur_total, buy_prev_total)*100:+.2f}%",
                "状态": get_status(calc_ratio(buy_cur_total, buy_prev_total))
            },
            {
                "指标": "下载转化率",
                f"{date_info['cur_label']}": f"{dl_c_rate*100:.2f}%",
                f"{date_info['prev_label']}": f"{dl_p_rate*100:.2f}%",
                "环比": f"{(dl_c_rate - dl_p_rate)*100:+.2f}%",
                "状态": get_status(dl_c_rate - dl_p_rate)
            },
            {
                "指标": "购买转化率",
                f"{date_info['cur_label']}": f"{buy_c_rate*100:.2f}%",
                f"{date_info['prev_label']}": f"{buy_p_rate*100:.2f}%",
                "环比": f"{(buy_c_rate - buy_p_rate)*100:+.2f}%",
                "状态": get_status(buy_c_rate - buy_p_rate)
            }
        ]
        overview_df = pd.DataFrame(overview_rows)

        site_col1, site_col2 = st.columns(2)
        site_col1.info(f"🌐 **负责站点**：`{site_name}`")
        site_col2.warning(f"⚡ **{ip_label}**：`{daily_ip:,}`")

        st.dataframe(overview_df, use_container_width=True, hide_index=True)
        st.markdown("---")

        # ================= 业务明细 Tab 看板 =================
        tab1, tab2, tab3 = st.tabs(["📥 客户端下载量看板", "💰 购买量看板", "🌐 渠道流量看板"])

        with tab1:
            dl_df = pd.DataFrame([
                {"产品包": k, f"本期 ({date_info['cur_label']})": v["cur"], f"上期 ({date_info['prev_label']})": v["prev"]}
                for k, v in dl_data.items()
            ])
            dl_df["差额"] = dl_df.iloc[:, 1] - dl_df.iloc[:, 2]
            dl_df["环比变动"] = dl_df.apply(lambda r: f"{(r['差额'] / r.iloc[2] * 100):.2f}%" if r.iloc[2] != 0 else "-", axis=1)

            chart_type = st.radio("展现方式：", ["柱状对比图", "折线走势图", "纯数据表格"], horizontal=True, key="dl_view")
            if chart_type == "柱状对比图":
                fig = px.bar(dl_df, x="产品包", y=[dl_df.columns[1], dl_df.columns[2]], barmode="group", title="各客户端下载量对比", color_discrete_sequence=["#1976D2", "#90CAF9"])
                st.plotly_chart(fig, use_container_width=True)
            elif chart_type == "折线走势图":
                fig = px.line(dl_df, x="产品包", y=[dl_df.columns[1], dl_df.columns[2]], markers=True, title="各客户端下载量走势对比", color_discrete_sequence=["#1976D2", "#FFA726"])
                st.plotly_chart(fig, use_container_width=True)
            else:
                st.dataframe(dl_df, use_container_width=True)

        with tab2:
            plan_choice = st.selectbox("选择套餐周期：", ["Lifetime", "12M", "3M", "1M"])
            target_plan_data = buy_data.get(plan_choice.lower() if plan_choice != "12M" else "12M", {})
            buy_df = pd.DataFrame([
                {"产品类型": k, f"本期 ({date_info['cur_label']})": v["cur"], f"上期 ({date_info['prev_label']})": v["prev"]}
                for k, v in target_plan_data.items()
            ])
            if not buy_df.empty:
                buy_df["差额"] = buy_df.iloc[:, 1] - buy_df.iloc[:, 2]
                buy_df["环比变动"] = buy_df.apply(lambda r: f"{(r['差额'] / r.iloc[2] * 100):.2f}%" if r.iloc[2] != 0 else "-", axis=1)

                buy_chart_type = st.radio("展现方式：", ["柱状对比图", "折线走势图", "纯数据表格"], horizontal=True, key="buy_view")
                if buy_chart_type == "柱状对比图":
                    fig = px.bar(buy_df, x="产品类型", y=[buy_df.columns[1], buy_df.columns[2]], barmode="group", title=f"{plan_choice} 购买转化量对比", color_discrete_sequence=["#4CAF50", "#A5D6A7"])
                    st.plotly_chart(fig, use_container_width=True)
                elif buy_chart_type == "折线走势图":
                    fig = px.line(buy_df, x="产品类型", y=[buy_df.columns[1], buy_df.columns[2]], markers=True, title=f"{plan_choice} 购买转化量对比", color_discrete_sequence=["#4CAF50", "#FFB74D"])
                    st.plotly_chart(fig, use_container_width=True)
                else:
                    st.dataframe(buy_df, use_container_width=True)
            else:
                st.info(f"暂无 {plan_choice} 维度的购买数据")

        with tab3:
            if tf_data:
                channels = list(set(list(tf_data["cur"].keys()) + list(tf_data["prev"].keys())))
                tf_df = pd.DataFrame([
                    {"渠道": c, f"本期 ({date_info['cur_label']})": tf_data["cur"].get(c, 0), f"上期 ({date_info['prev_label']})": tf_data["prev"].get(c, 0)}
                    for c in channels
                ])
                tf_df["差额"] = tf_df.iloc[:, 1] - tf_df.iloc[:, 2]
                tf_df["环比变动"] = tf_df.apply(lambda r: f"{(r['差额'] / r.iloc[2] * 100):.2f}%" if r.iloc[2] != 0 else "-", axis=1)

                tf_chart_type = st.radio("展现方式：", ["柱状对比图", "折线走势图", "纯数据表格"], horizontal=True, key="tf_view")
                if tf_chart_type == "柱状对比图":
                    fig = px.bar(tf_df, x="渠道", y=[tf_df.columns[1], tf_df.columns[2]], barmode="group", title="各渠道用户获取量对比", color_discrete_sequence=["#9C27B0", "#CE93D8"])
                    st.plotly_chart(fig, use_container_width=True)
                elif tf_chart_type == "折线走势图":
                    fig = px.line(tf_df, x="渠道", y=[tf_df.columns[1], tf_df.columns[2]], markers=True, title="各渠道用户获取量对比", color_discrete_sequence=["#9C27B0", "#FF7043"])
                    st.plotly_chart(fig, use_container_width=True)
                else:
                    st.dataframe(tf_df, use_container_width=True)
            else:
                st.info("未检测到《用户获取情况》CSV 文件，暂不展示渠道数据。")
else:
    st.info("👈 请在左侧侧边栏上传你的 CSV 文件以启动数据看板。")
